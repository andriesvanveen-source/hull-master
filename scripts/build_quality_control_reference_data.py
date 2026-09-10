import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "public" / "quality-control" / "Quality Control Reference Workbook.xlsx"
C2026_WORKBOOK = ROOT / "public" / "quality-control" / "reference-source" / "C2026 QC audit sheet.xlsx"
OUTPUT_JSON = ROOT / "public" / "quality-control" / "reference-audits.json"
OUTPUT_SQL = ROOT / "supabase" / "quality-control-reference-data.sql"
OUTPUT_SQL_PATTERN = "quality-control-reference-data-{model}.sql"
C2026_OUTPUT_SQL = ROOT / "supabase" / "quality-control-reference-data-C2026.sql"
DISCIPLINES = {1: "Gelcoat", 2: "Flowcoat", 3: "Joinery/Carp", 4: "Deckfitting", 5: "Plumbing", 6: "Mechanical", 7: "Electrical", 8: "Perspex/Windows", 9: "Spray Painting", 10: "Cleaning"}
STAMP = "2026-09-10T00:00:00+00:00"


def clean(value):
    return re.sub(r"\s+", " ", str(value or "").replace("�", "")).strip()


def canonical_area(value):
    original = clean(value)
    key = re.sub(r"[^a-z0-9]+", " ", original.lower()).strip()
    key = re.sub(r"\b(cont|continued)\b", "", key).strip()
    key = re.sub(r"\s+(1|2|a|b)$", "", key).strip()
    key = key.replace("starboard", "stbd").replace("stb ", "stbd ").replace("sth ", "stbd ")
    key = key.replace("engine rm", "engine room").replace("eng rm", "engine room")
    mappings = [
        (r"^aft cockpit", "Aft Cockpit"), (r"^(fwd|forward) cockpit.*fore ?deck|^fwd cockpit$", "Fwd Cockpit & Foredeck"),
        (r"^fore ?deck$", "Fwd Cockpit & Foredeck"), (r"^fore ?peak$|^port fore ?peak$", "Forepeak"),
        (r"^fly ?bridge$", "Flybridge"), (r"^saloon", "Saloon"), (r"^galley$", "Galley"),
        (r"^port aft cabin", "Port Aft Cabin"), (r"^port aft heads", "Port Aft Heads"),
        (r"^port fwd cabin", "Port Fwd Cabin"), (r"^port fwd heads", "Port Fwd Heads"),
        (r"^port mid(ship)? cabin", "Port Mid Cabin"), (r"^port mid(ship)? heads", "Port Mid Heads"),
        (r"^port passage", "Port Passage"), (r"^port side deck", "Port Side Deck"),
        (r"^port hull", "Port Hull"), (r"^port (engine room|engine)$", "Port Engine Room"),
        (r"^port heads$", "Port Heads"), (r"^port crew quarter", "Crew Quarters"),
        (r"^stbd aft cabin", "Stbd Aft Cabin"), (r"^stbd aft heads", "Stbd Aft Heads"),
        (r"^stbd fwd cabin", "Stbd Fwd Cabin"), (r"^stbd fwd heads", "Stbd Fwd Heads"),
        (r"^stbd (engine room|engine)$", "Stbd Engine Room"), (r"^stbd hull", "Stbd Hull"),
        (r"^stbd side deck", "Stbd Side Deck"), (r"^stbd passage$|^stbd fwd passage$", "Stbd Passage"),
        (r"^stbd (study|study passage|passage and study|cabin and passage)", "Stbd Study & Passage"),
        (r"^stbd ov cabin", "Stbd Owner's Cabin"), (r"^stbd ov (heads|ext heads)|^ov heads$", "Owner's Heads"),
        (r"^stbd ov study", "Owner's Study"), (r"^study$", "Study"), (r"^crew quarter", "Crew Quarters"),
    ]
    for pattern, label in mappings:
        if re.search(pattern, key):
            return label
    return original.title() or "Unassigned"


def is_heading(row):
    return not isinstance(row[2], (int, float)) and not row[3] and not row[4]


def heading_candidate(row):
    candidate = clean(row[6]) or clean(row[1])
    if not candidate or re.search(r"^(description|boat no|area|qc|date|column|exported)", candidate, re.I):
        return ""
    return candidate


def sql_text(value):
    return "'" + str(value or "").replace("'", "''") + "'"


def parse_audit(sheet, name, defect_id_prefix=None):
    current_area = "Unassigned"
    areas = []
    defects = []
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
        row = list(values) + [None] * 12
        if is_heading(row):
            candidate = heading_candidate(row)
            if candidate:
                current_area = canonical_area(candidate)
            continue
        code = row[2]
        if not isinstance(code, (int, float)) or int(code) not in DISCIPLINES:
            continue
        area_cell = clean(row[1])
        if area_cell and area_cell.lower() != "concern":
            current_area = canonical_area(area_cell)
        area = canonical_area(current_area)
        if area not in areas:
            areas.append(area)
        item = clean(row[3]) or "Unspecified item"
        failure = clean(row[4]) or "Defect"
        description = clean(row[6]) or clean(row[5]) or f"{item} — {failure}"
        defects.append({
            "id": f"{defect_id_prefix or f'qc-source-{name.lower()}'}-r{row_number}",
            "area": area, "item": item, "failure": failure, "description": description, "code": int(code),
            "discipline": DISCIPLINES[int(code)], "concern": area_cell.lower() == "concern",
            "repairedBy": clean(row[7]), "repairedDate": clean(row[8]),
            "teamLeaderCheck": clean(row[9]), "qcRwk": clean(row[10]), "qcAcc": clean(row[11]),
            "createdAt": STAMP, "updatedAt": STAMP,
        })
    return {
        "id": f"qc-reference-{name.lower()}", "name": name, "model": name[:2], "areas": areas,
        "areaInspectors": {}, "completedAreas": [], "defects": defects,
        "createdAt": STAMP, "updatedAt": STAMP,
    }


def build():
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    audits = [parse_audit(workbook[sheet_name], sheet_name[:-6]) for sheet_name in workbook.sheetnames if sheet_name.endswith(" Audit")]

    def sql_for_audits(selected_audits, label):
        lines = [f"-- Historical Quality Control import for {label}.", "-- Idempotent and non-destructive: existing boat, area and defect records are not overwritten.", "begin;"]
        for audit in selected_audits:
            lines.append(f"insert into public.quality_control_boats (id,name,model) values ({sql_text(audit['id'])},{sql_text(audit['name'])},{sql_text(audit['model'])}) on conflict (id) do nothing;")
            for index, area in enumerate(audit["areas"]):
                lines.append(f"insert into public.quality_control_areas (boat_id,area_name,sort_order) values ({sql_text(audit['id'])},{sql_text(area)},{index}) on conflict (boat_id,area_name) do nothing;")
            for start in range(0, len(audit["defects"]), 200):
                rows = []
                for defect in audit["defects"][start:start + 200]:
                    rows.append("(" + ",".join([
                        sql_text(defect["id"]), sql_text(audit["id"]), sql_text(defect["area"]), sql_text(defect["item"]),
                        sql_text(defect["failure"]), sql_text(defect["description"]), str(defect["code"]),
                        sql_text(defect["discipline"]), "true" if defect["concern"] else "false",
                        sql_text(defect["repairedBy"]), sql_text(defect["repairedDate"]), sql_text(defect["teamLeaderCheck"]),
                        sql_text(defect["qcRwk"]), sql_text(defect["qcAcc"])
                    ]) + ")")
                lines.append("insert into public.quality_control_defects (id,boat_id,area_name,item,failure_mode,description,code,discipline,concern,repaired_by,repaired_date,team_leader_check,qc_rwk,qc_acc) values\n  " + ",\n  ".join(rows) + "\non conflict (id) do nothing;")
        lines.extend(["commit;", "select count(*) as imported_quality_control_defects from public.quality_control_defects;"])
        return "\n".join(lines) + "\n"

    OUTPUT_SQL.write_text("-- This import was split to fit the Supabase SQL Editor.\n-- Run quality-control-reference-data-B5.sql, B8.sql, B9.sql, C1.sql, C2.sql and C5.sql instead.\n", encoding="utf-8")
    for model in ["B5", "B8", "B9", "C1", "C2", "C5"]:
        selected = [audit for audit in audits if audit["model"] == model]
        (OUTPUT_SQL.parent / OUTPUT_SQL_PATTERN.format(model=model)).write_text(sql_for_audits(selected, f"{model} boats"), encoding="utf-8")

    c2026_workbook = openpyxl.load_workbook(C2026_WORKBOOK, read_only=True, data_only=True)
    c2026_supplement = parse_audit(c2026_workbook["Audit sheet"], "C2026", "qc-source-c2026-qc3")
    c2026_audit = next(audit for audit in audits if audit["name"] == "C2026")
    c2026_audit["areas"] = list(dict.fromkeys([*c2026_audit["areas"], *c2026_supplement["areas"]]))
    c2026_audit["defects"].extend(c2026_supplement["defects"])
    c2026_audit["updatedAt"] = STAMP
    OUTPUT_JSON.write_text(json.dumps(audits, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    C2026_OUTPUT_SQL.write_text(sql_for_audits([c2026_supplement], "C2026 supplemental audit"), encoding="utf-8")
    print(f"Generated {len(audits)} audits and {sum(len(a['defects']) for a in audits)} defects, including {len(c2026_supplement['defects'])} supplemental C2026 defects")


if __name__ == "__main__":
    build()
